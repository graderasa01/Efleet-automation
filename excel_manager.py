from __future__ import annotations

from pathlib import Path
from threading import Lock
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook, Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .query_engine import AUDIT_COLUMNS, build_audit_queries, default_row_values, resolve_query


class ExcelManager:
    """Safe xlsx updater for Codespaces/server usage.

    This uses openpyxl because Codespaces cannot control a live desktop Excel app.
    It preserves normal workbook content and formulas much better than pandas df.to_excel.
    """

    def __init__(self, file_path: str, sheet_name: str = "DATA", search_column_name: str = "HireSlipNo", search_column_index_1_based: int = 2):
        self.file_path = Path(file_path)
        self.sheet_name = sheet_name
        self.search_column_name = search_column_name
        self.search_column_index_1_based = search_column_index_1_based
        self.lock = Lock()
        self.wb = None
        self.ws: Optional[Worksheet] = None
        self.headers: List[str] = []
        self.header_to_col: Dict[str, int] = {}

    def open(self) -> None:
        with self.lock:
            if not self.file_path.exists():
                raise FileNotFoundError(f"Excel file not found: {self.file_path}. Put your file here or update config.")
            self.wb = load_workbook(self.file_path)
            if self.sheet_name in self.wb.sheetnames:
                self.ws = self.wb[self.sheet_name]
            else:
                self.ws = self.wb.active
            self._load_headers_locked()
            self._ensure_columns_locked(["Bot Status", "Bot Error", "Audit Queries", "Advance department remark"])
            self.save_locked()

    def _load_headers_locked(self) -> None:
        assert self.ws is not None
        self.headers = []
        self.header_to_col = {}
        for cell in self.ws[1]:
            text = "" if cell.value is None else str(cell.value).strip()
            self.headers.append(text)
            if text and text not in self.header_to_col:
                self.header_to_col[text] = cell.column

    def _ensure_columns_locked(self, names: List[str]) -> None:
        assert self.ws is not None
        for name in names:
            if name not in self.header_to_col:
                col = self.ws.max_column + 1
                self.ws.cell(row=1, column=col, value=name)
                self.header_to_col[name] = col
                self.headers.append(name)

    def save_locked(self) -> None:
        assert self.wb is not None
        self.file_path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(self.file_path)

    def save(self) -> None:
        with self.lock:
            self.save_locked()

    def _resolve_search_col_locked(self) -> int:
        assert self.ws is not None
        # Prefer exact name if present.
        if self.search_column_name in self.header_to_col:
            return self.header_to_col[self.search_column_name]
        # Fallback: first header that contains hireslip.
        for header, col in self.header_to_col.items():
            if "hireslip" in header.lower().replace(" ", ""):
                return col
        # Final fallback: configured numeric index.
        return int(self.search_column_index_1_based)

    def iter_data_rows(self, start_excel_row: int = 2) -> List[Tuple[int, str]]:
        with self.lock:
            assert self.ws is not None
            search_col = self._resolve_search_col_locked()
            rows: List[Tuple[int, str]] = []
            for row_idx in range(start_excel_row, self.ws.max_row + 1):
                value = self.ws.cell(row=row_idx, column=search_col).value
                hm = "" if value is None else str(value).strip()
                if hm and hm.lower() != "nan":
                    rows.append((row_idx, hm))
            return rows

    def get_row_dict(self, excel_row: int) -> Dict[str, Any]:
        with self.lock:
            assert self.ws is not None
            self._load_headers_locked()
            data: Dict[str, Any] = {}
            for header, col in self.header_to_col.items():
                data[header] = self.ws.cell(row=excel_row, column=col).value
            return data

    def set_row_status(self, excel_row: int, status: str, error: str = "") -> None:
        with self.lock:
            assert self.ws is not None
            self._ensure_columns_locked(["Bot Status", "Bot Error"])
            self.ws.cell(row=excel_row, column=self.header_to_col["Bot Status"], value=status)
            self.ws.cell(row=excel_row, column=self.header_to_col["Bot Error"], value=error)
            self.save_locked()

    def apply_defaults_to_row(self, excel_row: int) -> None:
        with self.lock:
            assert self.ws is not None
            self._load_headers_locked()
            self._ensure_columns_locked(AUDIT_COLUMNS + ["Audit Queries"])
            defaults = default_row_values()
            for col_name, value in defaults.items():
                col = self.header_to_col[col_name]
                current = self.ws.cell(row=excel_row, column=col).value
                if current is None or str(current).strip() == "":
                    self.ws.cell(row=excel_row, column=col, value=value)
            self._rebuild_audit_queries_locked(excel_row)
            self.save_locked()

    def apply_query(self, excel_row: int, column: str, query_no: str = "", custom_text: str = "") -> str:
        with self.lock:
            assert self.ws is not None
            self._load_headers_locked()
            self._ensure_columns_locked([column, "Audit Queries"])
            text = resolve_query(column, query_no, custom_text)
            self.ws.cell(row=excel_row, column=self.header_to_col[column], value=text)
            audit = self._rebuild_audit_queries_locked(excel_row)
            self.save_locked()
            return audit

    def write_cell_by_header(self, excel_row: int, header: str, value: Any) -> None:
        with self.lock:
            assert self.ws is not None
            self._ensure_columns_locked([header])
            self.ws.cell(row=excel_row, column=self.header_to_col[header], value=value)
            self.save_locked()

    def _rebuild_audit_queries_locked(self, excel_row: int) -> str:
        assert self.ws is not None
        self._ensure_columns_locked(AUDIT_COLUMNS + ["Audit Queries"])
        row_values = {col: self.ws.cell(row=excel_row, column=self.header_to_col[col]).value for col in AUDIT_COLUMNS}
        audit = build_audit_queries(row_values)
        self.ws.cell(row=excel_row, column=self.header_to_col["Audit Queries"], value=audit)
        return audit
